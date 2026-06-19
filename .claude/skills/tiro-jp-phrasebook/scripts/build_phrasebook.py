# -*- coding: utf-8 -*-
"""
Tiro 미팅 → 비즈니스 일본어 표현집(xlsx) 생성기.

레이아웃 규칙:
  - 열: No. | カテゴリ (분류) | 文章 (日本語 / 한국어)
  - 한국어 번역은 일본어 셀 '바로 아래 행'에 같은 열로 배치
  - No./카테고리 셀은 두 행에 걸쳐 병합
  - 일본어 행: 굵은 남색(흰 배경) / 한국어 행: 초록 이탤릭(연한 파란 배경)

사용법:
  1) 아래 ROWS를 (카테고리, 일본어, 한국어) 튜플로 채운다.
  2) python3 build_phrasebook.py  (또는 build_workbook(ROWS, out_path) 호출)
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def build_workbook(rows, out_path, sheet_title="ビジネス日本語表現集"):
    """rows: list of (category, japanese, korean). out_path: .xlsx 경로."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title

    headers = ["No.", "カテゴリ (분류)", "文章 (日本語 / 한국어)"]
    ws.append(headers)

    navy = PatternFill("solid", fgColor="2F5496")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    jp_font = Font(bold=True, size=11, color="1F3864")
    ko_font = Font(italic=True, size=10, color="375623")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    jp_fill = PatternFill("solid", fgColor="FFFFFF")
    ko_fill = PatternFill("solid", fgColor="EDF2FB")

    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col)
        c.fill = navy
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border

    r = 2
    for i, (cat, jp, ko) in enumerate(rows, 1):
        jp_row, ko_row = r, r + 1
        ws.cell(row=jp_row, column=1, value=i)
        ws.cell(row=jp_row, column=2, value=cat)
        ws.cell(row=jp_row, column=3, value=jp)
        ws.cell(row=ko_row, column=3, value=ko)

        ws.merge_cells(start_row=jp_row, start_column=1, end_row=ko_row, end_column=1)
        ws.merge_cells(start_row=jp_row, start_column=2, end_row=ko_row, end_column=2)

        ws.cell(row=jp_row, column=3).font = jp_font
        ws.cell(row=ko_row, column=3).font = ko_font
        for rr in (jp_row, ko_row):
            ws.cell(row=rr, column=3).fill = jp_fill if rr == jp_row else ko_fill
            for col in range(1, 4):
                cell = ws.cell(row=rr, column=col)
                cell.border = border
                cell.alignment = Alignment(
                    vertical="center", wrap_text=True,
                    horizontal="center" if col in (1, 2) else "left",
                )
        r += 2

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 90
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"

    wb.save(out_path)
    return out_path


# ── 미팅에서 추출·정정한 문장으로 교체해 사용 ──────────────────
# (카테고리, 비즈니스 일본어, 한국어 번역)
ROWS = [
    ("会議の開始・挨拶", "本日はお忙しい中、お時間をいただきありがとうございます。",
     "오늘 바쁘신 와중에 시간 내주셔서 감사합니다."),
    ("中継事故対応", "万が一、中継事故が発生した場合の対応フローを、事前に共有させていただきます。",
     "만일 중계 사고가 발생했을 경우의 대응 플로우를 사전에 공유드리겠습니다."),
    # … 미팅 내용에 맞춰 50~100문장으로 채운다 …
]


if __name__ == "__main__":
    out = build_workbook(ROWS, "cosme_business_japanese.xlsx")
    print("Saved:", out, "| sentences:", len(ROWS))
