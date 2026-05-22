import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

data = [
    # brand_jp, brand_en, category, country, instagram, twitter, notes
    ("COSRX", "COSRX", "スキンケア", "韓国", "https://www.instagram.com/cosrx_jp", "https://x.com/COSRX_JAPAN", ""),
    ("ラネージュ", "Laneige", "スキンケア/メイク", "韓国", "https://www.instagram.com/laneige_jp", "https://x.com/Laneige_japan", ""),
    ("ロムアンド", "rom&nd", "メイク", "韓国", "https://www.instagram.com/romand_jp", "https://x.com/romandyou", ""),
    ("ティルティル", "TIRTIR", "スキンケア/メイク", "韓国", "https://www.instagram.com/tirtir_jp_official", "https://x.com/tirtir_jp", ""),
    ("アヌア", "Anua", "スキンケア", "韓国", "https://www.instagram.com/anua.jp", "https://x.com/anua_official", ""),
    ("ダルバ", "d'Alba", "スキンケア", "韓国", "https://www.instagram.com/dalba_japan", "https://x.com/dAlba_japan", ""),
    ("VTコスメ", "VT Cosmetics", "スキンケア", "韓国", "https://www.instagram.com/vtcosmetics_japan", "https://x.com/vtcosmetics_jp", ""),
    ("SKIN1004", "SKIN1004", "スキンケア", "韓国", "https://www.instagram.com/skin1004_japan", "https://x.com/skin1004_japan", ""),
    ("メディキューブ", "Medicube", "スキンケア", "韓国", "https://www.instagram.com/medicube_officialjapan", "https://x.com/medicube_japan", ""),
    ("I'm From", "I'm From", "スキンケア", "韓国", "https://www.instagram.com/imfrom_jp", "https://x.com/imfrom_jp", ""),
    ("バニラコ", "Banila Co", "メイク", "韓国", "https://www.instagram.com/banilaco_japan", "https://x.com/Banilaco_Japan", ""),
    ("バイオヒールボ", "Bioheal BOH", "スキンケア", "韓国", "https://www.instagram.com/bioheal.boh_japan", "https://x.com/BioHealBohJapan", ""),
    ("ペリペラ", "Peripera", "メイク", "韓国", "https://www.instagram.com/periperajapan", "https://x.com/periperajapan", ""),
    ("スキンアンドラブ", "Skin&Lab", "スキンケア", "韓国", "https://www.instagram.com/skinnlab_japan", "https://x.com/skinnlab_japan", ""),
    ("FLOWER KNOWS", "Flower Knows", "メイク", "中国", "https://www.instagram.com/flowerknows_jp", "https://x.com/flowerknows_jp", ""),
    ("manyo", "Manyo Factory", "スキンケア", "韓国", "https://www.instagram.com/manyo.japan", "https://x.com/manyojapan", ""),
    ("LUNA", "LUNA", "メイク", "韓国", "https://www.instagram.com/luna_makeup_jp", "https://x.com/luna_makeup_jp", ""),
    ("ISOI", "ISOI", "スキンケア", "韓国", "https://www.instagram.com/isoi_japan", "https://x.com/isoi_jp", ""),
    ("TFIT", "TFIT", "スキンケア", "韓国", "https://www.instagram.com/tfit.japan", "https://x.com/TFIT_JP", ""),
    ("BLANC", "Blanc Nature", "スキンケア", "韓国", "https://www.instagram.com/blancnature_jp", "https://x.com/blanc_jp", ""),
    # Instagram only
    ("エイプリルスキン", "April Skin", "スキンケア/メイク", "韓国", "https://www.instagram.com/aprilskin_officialjapan", "未確認", ""),
    ("トゥークールフォースクール", "too cool for school", "メイク", "韓国", "https://www.instagram.com/toocoolforschool_official_jp", "未確認", ""),
    ("センテリアン24", "Centellian24", "スキンケア", "韓国", "https://www.instagram.com/centellian24_japan", "未確認", ""),
    ("カラーグラム", "Colorgram", "メイク", "韓国", "https://www.instagram.com/colorgram_jp", "未確認", ""),
    ("Real Barrier", "Real Barrier", "スキンケア", "韓国", "https://www.instagram.com/realbarrier_jp_official", "未確認", ""),
    ("ink.", "ink.", "メイク", "韓国", "https://www.instagram.com/ink.129", "未確認", ""),
    ("PERFECT DIARY", "Perfect Diary", "メイク", "中国", "https://www.instagram.com/perfectdiary_japan", "未確認", ""),
    # Twitter only
    ("ミルクタッチ", "Milk Touch", "メイク", "韓国", "未確認", "https://x.com/milktouch_jp", ""),
    # No confirmed Japan SNS
    ("ENTROPY MAKEUP", "Entropy Makeup", "メイク", "韓国", "未確認", "未確認", ""),
    ("Slow Humming", "Slow Humming", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("Furriky", "Furriky", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("UIQ", "UIQ", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("LEPOREM", "Leporem", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("CHASIN' RABBITS", "Chasin' Rabbits", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("HERRSAINE", "Herrsaine", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("23yearsold", "23yearsold", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("NESH", "NESH", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("NDP", "NDP", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("リーダース", "Leaders", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("Kind Step", "Kind Step", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("Hebestem", "Hebestem", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("マリエラン", "Mariela'n", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("MAWLAB", "MAWLAB", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("FINGER SUIT", "Finger Suit", "メイク", "韓国", "未確認", "未確認", ""),
    ("VELUS", "VELUS", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("TONE fit SUN", "TONE fit SUN", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("OUKEYA", "OUKEYA", "スキンケア", "中国", "未確認", "未確認", ""),
    ("IZEZE", "IZEZE", "メイク", "中国", "未確認", "未確認", ""),
    ("Barle", "Barle", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("Brilliet", "Brilliet", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("efilow", "efilow", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("GIK", "GIK", "メイク", "韓国", "未確認", "未確認", ""),
    ("KISOCARE", "KISOCARE", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("2aN", "2aN", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("soaddicted", "soaddicted", "メイク", "韓国", "未確認", "未確認", ""),
    ("XIXI", "XIXI", "メイク", "中国", "未確認", "未確認", ""),
    ("It's Skin", "It's Skin", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("EIoM", "EIoM", "メイク", "韓国", "未確認", "未確認", ""),
    ("celimax", "celimax", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("athe", "athe", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("TONE28", "TONE28", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("PHYTOVER", "PHYTOVER", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("fremena", "fremena", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("lily eve", "lily eve", "メイク", "韓国", "未確認", "未確認", ""),
    ("isntree", "isntree", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("ミュード", "Myude", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("LIALUSTER", "LIALUSTER", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("ヌニ", "NUNI", "メイク", "韓国", "未確認", "未確認", ""),
    ("OIAD", "OIAD", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("インセルダーム", "Inselderm", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("エリクシール", "Elixir", "スキンケア", "日本", "未確認", "未確認", "資生堂ブランド"),
    ("FRANZ", "FRANZ", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("リブコイ", "Livkoi", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("セラディックス", "Celadix", "スキンケア", "韓国", "未確認", "未確認", ""),
    ("fwee", "fwee", "メイク", "韓国", "未確認", "未確認", ""),
    ("JUDYDOLL", "Judydoll", "メイク", "中国", "未確認", "未確認", ""),
    ("NIKI PITA", "NIKI PITA", "スキンケア", "韓国", "未確認", "未確認", ""),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Qoo10 化粧品ブランドSNS"

# Header
headers = ["ブランド名（日本語）", "ブランド名（英語）", "カテゴリ", "原産国", "Instagram Japan", "Twitter/X Japan", "備考"]
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF", size=11)
border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = border

# Color fills for rows
green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
light_blue = PatternFill(start_color="DEEAF1", end_color="DEEAF1", fill_type="solid")

for row_idx, row_data in enumerate(data, 2):
    brand_jp, brand_en, category, country, instagram, twitter, notes = row_data
    has_ig = instagram != "未確認"
    has_tw = twitter != "未確認"
    if has_ig and has_tw:
        row_fill = green_fill
    elif has_ig or has_tw:
        row_fill = yellow_fill
    else:
        row_fill = red_fill

    values = [brand_jp, brand_en, category, country, instagram, twitter, notes]
    for col_idx, value in enumerate(values, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.fill = row_fill
        cell.border = border
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        if col_idx in (5, 6) and value != "未確認":
            cell.font = Font(color="0563C1", underline="single")
        else:
            cell.font = Font(size=10)

# Column widths
col_widths = [22, 22, 18, 8, 45, 45, 18]
for i, width in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = width

ws.row_dimensions[1].height = 30

# Legend sheet
ws2 = wb.create_sheet("凡例")
ws2.column_dimensions['A'].width = 30
ws2.column_dimensions['B'].width = 40
legend_data = [
    ("色", "意味"),
    ("緑", "Instagram & Twitter/X 両方確認済み"),
    ("黄", "Instagram または Twitter/X のみ確認済み"),
    ("赤/オレンジ", "日本公式SNS 未確認"),
]
legend_fills = [
    PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid"),
    green_fill,
    yellow_fill,
    red_fill,
]
legend_fonts = [
    Font(bold=True, color="FFFFFF"),
    Font(size=10),
    Font(size=10),
    Font(size=10),
]
for r, (label, desc) in enumerate(legend_data, 1):
    c1 = ws2.cell(row=r, column=1, value=label)
    c2 = ws2.cell(row=r, column=2, value=desc)
    c1.fill = legend_fills[r-1]
    c2.fill = legend_fills[r-1]
    c1.font = legend_fonts[r-1]
    c2.font = legend_fonts[r-1]
    c1.border = border
    c2.border = border
    c1.alignment = Alignment(horizontal='center', vertical='center')
    c2.alignment = Alignment(vertical='center')
    ws2.row_dimensions[r].height = 22

output_path = "/home/user/mypage/qoo10_beauty_brands_sns.xlsx"
wb.save(output_path)
print(f"Saved: {output_path}")
print(f"Total brands: {len(data)}")
confirmed_both = sum(1 for d in data if d[4] != "未確認" and d[5] != "未確認")
confirmed_one = sum(1 for d in data if (d[4] != "未確認") != (d[5] != "未確認"))
print(f"Both SNS confirmed: {confirmed_both}")
print(f"One SNS confirmed: {confirmed_one}")
print(f"No SNS confirmed: {len(data) - confirmed_both - confirmed_one}")
