#!/usr/bin/env python3
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import copy

W, H = Inches(13.33), Inches(7.5)   # 16:9 와이드

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H

BLANK = prs.slide_layouts[6]   # 완전 빈 슬라이드

# ── 색상 상수 ──────────────────────────────────────────────────────────────
NAVY   = RGBColor(0x1A, 0x23, 0x7E)
D1     = RGBColor(0xC0, 0x39, 0x2B)
D2     = RGBColor(0x15, 0x65, 0xC0)
D3     = RGBColor(0x2E, 0x7D, 0x32)
D4     = RGBColor(0xE6, 0x51, 0x00)
STAY   = RGBColor(0x7B, 0x1F, 0xA2)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LGRAY  = RGBColor(0xF0, 0xF4, 0xF8)
DGRAY  = RGBColor(0x37, 0x47, 0x4F)
SEA    = RGBColor(0xA8, 0xD8, 0xEA)

def add_rect(slide, l, t, w, h, fill, alpha=None):
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    return shape

def add_text(slide, text, l, t, w, h, size, bold=False, color=RGBColor(0,0,0),
             align=PP_ALIGN.LEFT, wrap=True):
    txb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    txb.word_wrap = wrap
    tf = txb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = 'Malgun Gothic'
    return txb

def add_multiline(slide, lines, l, t, w, h, size, bold=False,
                  color=RGBColor(0,0,0), align=PP_ALIGN.LEFT, line_spacing=1.15):
    from pptx.oxml.ns import qn
    from lxml import etree
    txb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    txb.word_wrap = True
    tf = txb.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = 'Malgun Gothic'
    return txb

# ════════════════════════════════════════════════════════════════════════════
# 슬라이드 1 — 타이틀
# ════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK)

# 배경 그라데이션 느낌 (두 레이어)
add_rect(slide, 0, 0, 13.33, 7.5, SEA)
add_rect(slide, 0, 0, 13.33, 3.8, NAVY)

# 지도 이미지
try:
    slide.shapes.add_picture('okinawa_travel_map.png',
                              Inches(7.5), Inches(0.15), Inches(5.6), Inches(7.2))
except:
    pass

# 타이틀 텍스트
add_text(slide, '오키나와', 0.5, 0.6, 6.5, 1.2, 52, bold=True, color=WHITE, align=PP_ALIGN.LEFT)
add_text(slide, '3박 4일 여행 코스', 0.5, 1.6, 6.5, 1.0, 36, bold=True, color=WHITE, align=PP_ALIGN.LEFT)
add_text(slide, '5월 26일(화) ~ 5월 29일(금)', 0.5, 2.55, 6.5, 0.6, 20, bold=False, color=RGBColor(0xB3,0xC8,0xF0), align=PP_ALIGN.LEFT)

# 항공 뱃지
add_rect(slide, 0.5, 3.3, 6.3, 0.65, D1)
add_text(slide, '✈  출발  5/26(화)  ICN → OKA  11:30–14:00  ZE631',
         0.65, 3.35, 6.0, 0.55, 15, bold=True, color=WHITE)
add_rect(slide, 0.5, 4.05, 6.3, 0.65, D4)
add_text(slide, '✈  귀국  5/29(금)  OKA → ICN  15:00–17:35  ZE632',
         0.65, 4.10, 6.0, 0.55, 15, bold=True, color=WHITE)

# 숙소 뱃지
add_rect(slide, 0.5, 4.9, 6.3, 0.65, STAY)
add_text(slide, '🏠  숙소  세소코섬 빌라  (3박)  Sesoko Island, Motobu',
         0.65, 4.95, 6.0, 0.55, 14, bold=True, color=WHITE)

# 예약번호
add_text(slide, '예약번호: VFZ8ME', 0.5, 6.8, 3.0, 0.5, 13, color=DGRAY)

# ════════════════════════════════════════════════════════════════════════════
# 슬라이드 2 — Day 1
# ════════════════════════════════════════════════════════════════════════════
def day_slide(prs, day_num, date_str, title_str, color, items):
    slide = prs.slides.add_slide(BLANK)
    add_rect(slide, 0, 0, 13.33, 7.5, LGRAY)

    # 좌측 색상 바
    add_rect(slide, 0, 0, 0.18, 7.5, color)

    # 헤더 바
    add_rect(slide, 0.18, 0, 13.15, 1.55, color)
    add_text(slide, day_num, 0.35, 0.08, 2.0, 0.7, 40, bold=True, color=WHITE)
    add_text(slide, date_str, 2.3, 0.10, 4.5, 0.65, 22, bold=True, color=RGBColor(0xFF,0xFF,0xFF))
    add_text(slide, title_str, 0.35, 0.78, 12.5, 0.7, 24, bold=True, color=WHITE)

    # 일정 테이블
    row_h = (7.5 - 1.75) / len(items)
    for i, (time_s, desc) in enumerate(items):
        y = 1.65 + i * row_h
        bg = RGBColor(0xFF,0xFF,0xFF) if i % 2 == 0 else RGBColor(0xF5,0xF5,0xF5)
        add_rect(slide, 0.18, y, 13.15, row_h - 0.04, bg)

        # 시간 칼럼
        add_rect(slide, 0.25, y + 0.06, 1.55, row_h - 0.18, color)
        add_text(slide, time_s, 0.28, y + 0.10, 1.45, row_h - 0.25,
                 19, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

        # 설명
        add_text(slide, desc, 2.05, y + 0.10, 11.1, row_h - 0.2,
                 20, bold=True, color=DGRAY)

    return slide

day_slide(prs, 'Day 1', '5월 26일 (화)', '나하 도착 → 세소코 정착', D1, [
    ('14:00', '나하 공항 도착  /  렌터카 수령'),
    ('14:30', '오키나와 자동차도로 북상  (약 1시간 30분)'),
    ('16:30', '세소코섬 빌라 체크인'),
    ('17:30', '세소코 비치 산책 & 석양 감상'),
    ('19:00', '모토부 현지 식당 저녁  (오키나와 소바 추천)'),
])

day_slide(prs, 'Day 2', '5월 27일 (수)', '북부 핵심 관광', D2, [
    ('08:30', '추라우미 수족관  (세계 최대급  /  빌라에서 10분)'),
    ('11:00', '오션 엑스포 공원 관람'),
    ('13:00', '비세 후쿠기 가로수길 근처 카페 점심'),
    ('14:00', '비세 후쿠기 가로수길  (수백 년 나무 터널)'),
    ('15:30', '나키진 성터  (유네스코 세계문화유산  /  바다 전망)'),
    ('18:00', '세소코 귀환  /  모토부 항구 해산물 저녁'),
])

day_slide(prs, 'Day 3', '5월 28일 (목)', '고우리섬 & 아메리칸 빌리지 쇼핑', D3, [
    ('09:00', '고우리 대교 드라이브  (에메랄드 바다 위 2km 다리)'),
    ('09:30', '티누 해변  (하트 바위 포토스팟)'),
    ('11:00', '고우리섬 점심  (고우리 새우 요리)'),
    ('13:00', '아메리칸 빌리지 쇼핑몰 이동  (약 1시간 30분)'),
    ('14:30', '아메리칸 빌리지 쇼핑 & 포토스팟'),
    ('18:00', '세소코 귀환'),
])

day_slide(prs, 'Day 4', '5월 29일 (금)', '국제거리 관광 → 귀국', D4, [
    ('08:00', '세소코 빌라 체크아웃  /  남하 시작'),
    ('10:00', '국제거리 쇼핑  (기념품)'),
    ('11:00', '마키시 공설시장 점심  (2층 해산물 정식)'),
    ('12:45', '렌터카 반납  &  나하 공항 이동'),
    ('13:00', '나하 공항 도착  /  탑승 수속'),
    ('15:00', 'OKA 출발 ✈  →  ICN 17:35 도착'),
])

# ════════════════════════════════════════════════════════════════════════════
# 슬라이드 6 — 숙소 + 꿀팁
# ════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK)
add_rect(slide, 0, 0, 13.33, 7.5, LGRAY)
add_rect(slide, 0, 0, 0.18, 7.5, STAY)
add_rect(slide, 0.18, 0, 13.15, 1.3, STAY)
add_text(slide, '숙소 정보  &  여행 꿀팁', 0.35, 0.25, 12.5, 0.85,
         32, bold=True, color=WHITE)

# 숙소 박스
add_rect(slide, 0.35, 1.45, 5.8, 2.6, WHITE)
add_rect(slide, 0.35, 1.45, 5.8, 0.55, STAY)
add_text(slide, '🏠  세소코섬 빌라  (숙소 · 3박)', 0.5, 1.48, 5.5, 0.5,
         17, bold=True, color=WHITE)
add_multiline(slide, [
    '주소: Okinawa, Motobu-cho, Sesoko 498-3',
    '예약번호: VFZ8ME',
    '비고: 세소코섬은 다리로 연결  /  렌터카 필수',
    '세소코 비치까지 도보 이동 가능',
], 0.5, 2.1, 5.5, 1.8, 15, color=DGRAY)

# 꿀팁 박스
add_rect(slide, 6.7, 1.45, 6.3, 2.6, WHITE)
add_rect(slide, 6.7, 1.45, 6.3, 0.55, RGBColor(0xF5,0x7F,0x17))
add_text(slide, '★  여행 꿀팁', 6.85, 1.48, 6.0, 0.5,
         17, bold=True, color=WHITE)
add_multiline(slide, [
    '• 렌터카: 공항 도착 즉시 수령 (북부 대중교통 불편)',
    '• 날씨: 5월 말은 장마 초입 — 우비/우산 필수',
    '• 추라우미 수족관: 오전 일찍 입장 추천',
    '• 고우리 새우: 섬 내 양식장 식당 강력 추천',
], 6.85, 2.1, 6.0, 1.8, 15, color=DGRAY)

# 항공편 요약
add_rect(slide, 0.35, 4.25, 12.65, 0.75, D1)
add_text(slide, '출발  5/26(화)  ICN → OKA  11:30 – 14:00  Eastar Jet ZE631  /  예약번호: VFZ8ME',
         0.55, 4.32, 12.2, 0.6, 16, bold=True, color=WHITE)
add_rect(slide, 0.35, 5.1, 12.65, 0.75, D4)
add_text(slide, '귀국  5/29(금)  OKA → ICN  15:00 – 17:35  Eastar Jet ZE632  /  예약번호: VFZ8ME',
         0.55, 5.17, 12.2, 0.6, 16, bold=True, color=WHITE)

# ════════════════════════════════════════════════════════════════════════════
# 슬라이드 7 — 지도
# ════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK)
add_rect(slide, 0, 0, 13.33, 7.5, LGRAY)
add_rect(slide, 0, 0, 13.33, 0.9, NAVY)
add_text(slide, '오키나와 여행 코스 지도', 0.4, 0.12, 12.5, 0.7,
         28, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
try:
    slide.shapes.add_picture('okinawa_travel_map.png',
                              Inches(0.5), Inches(1.0), Inches(12.33), Inches(6.3))
except:
    pass

out = 'okinawa_trip.pptx'
prs.save(out)
print(f"Saved: {out}")
